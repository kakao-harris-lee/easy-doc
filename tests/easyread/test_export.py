"""내보내기 산출물 생성 테스트.

여기서 보려는 것은 세 가지다: 자리표시자가 원문으로 온전히 복원되는지, docx·hwpx가
실제로 다시 열리는 파일인지(왕복), 그리고 한글 제목이 담긴 파일명 헤더가 규격을 지키는지.

hwpx 왕복은 **우리 ingest 추출기로** 확인한다 — 저장소에 한컴 오피스가 없어 실제 열람은
검증할 수 없고(그 한계는 export.py에 적혀 있다), 우리가 읽을 수 있는 구조로 조립됐는지가
지금 자동으로 지킬 수 있는 최대치다.
"""

import io
import zipfile
from urllib.parse import quote

import docx
import pytest

from app.easyread.export import (
    ExportFormat,
    content_disposition,
    export_filename,
    render_export,
    restore_placeholders,
)
from app.ingest.extractors import extract_text


def _paragraphs(content: bytes) -> list[str]:
    """생성한 docx를 다시 열어 문단 텍스트를 뽑는다."""
    return [paragraph.text for paragraph in docx.Document(io.BytesIO(content)).paragraphs]


# --- 자리표시자 복원 -------------------------------------------------------------


def test_자리표시자를_원문으로_되돌린다() -> None:
    text = "번호는 [[주민등록번호1]] 또는 [[카드번호1]]이에요."

    restored = restore_placeholders(
        text, {"[[주민등록번호1]]": "900101-1234567", "[[카드번호1]]": "1234-5678-9012-3456"}
    )

    assert restored == "번호는 900101-1234567 또는 1234-5678-9012-3456이에요."
    assert "[[" not in restored


def test_같은_자리표시자가_여러_번_나와도_모두_되돌린다() -> None:
    text = "[[주민등록번호1]]을 확인하세요. 다시 적으면 [[주민등록번호1]]입니다."

    restored = restore_placeholders(text, {"[[주민등록번호1]]": "900101-1234567"})

    assert restored.count("900101-1234567") == 2


def test_번호가_이어지는_자리표시자를_섞지_않는다() -> None:
    """[[주민등록번호1]] 치환이 [[주민등록번호11]]을 갉아먹으면 안 된다.

    닫는 괄호까지 통째로 보아야 한다.
    """
    text = "[[주민등록번호1]]과 [[주민등록번호11]]"

    restored = restore_placeholders(
        text, {"[[주민등록번호1]]": "900101-1111111", "[[주민등록번호11]]": "900101-2222222"}
    )

    assert restored == "900101-1111111과 900101-2222222"


def test_가린_항목이_없으면_본문이_그대로다() -> None:
    assert restore_placeholders("가린 것이 없는 글", {}) == "가린 것이 없는 글"


# --- docx ---------------------------------------------------------------------


def test_docx는_제목과_문단으로_구성된다() -> None:
    file = render_export(
        export_format=ExportFormat.DOCX,
        title="재난지원금 안내",
        body="첫 문단입니다.\n\n둘째 문단입니다.",
    )

    # 왕복: 우리가 만든 파일을 python-docx가 다시 열 수 있어야 한다.
    assert _paragraphs(file.content) == ["재난지원금 안내", "첫 문단입니다.", "둘째 문단입니다."]
    assert file.filename == "재난지원금 안내.docx"
    assert file.media_type.endswith("wordprocessingml.document")


def test_docx에_꼬리말을_붙이지_않는다() -> None:
    """담당자가 그대로 배포할 산출물이다 — "AI 초안" 같은 문구가 남으면 지워야 한다."""
    file = render_export(export_format=ExportFormat.DOCX, title="안내문", body="본문입니다.")

    assert _paragraphs(file.content) == ["안내문", "본문입니다."]


def test_문단_안의_줄바꿈은_유지된다() -> None:
    """빈 줄만 문단 경계다 — 목록처럼 줄만 바꾼 곳이 한 줄로 뭉치면 안 된다."""
    file = render_export(
        export_format=ExportFormat.DOCX, title="안내문", body="첫째 줄\n둘째 줄\n\n다음 문단"
    )

    assert _paragraphs(file.content) == ["안내문", "첫째 줄\n둘째 줄", "다음 문단"]


def test_빈_줄이_여러_개여도_빈_문단을_만들지_않는다() -> None:
    file = render_export(
        export_format=ExportFormat.DOCX, title="안내문", body="\n\n앞\n\n\n\n뒤\n\n"
    )

    assert _paragraphs(file.content) == ["안내문", "앞", "뒤"]


def test_제어문자가_섞여도_docx를_만든다() -> None:
    """XML은 탭·개행·복귀를 뺀 제어문자를 담지 못해 lxml이 ValueError를 던진다.

    제목은 저장 시점에 이미 걸러지지만 초안 본문은 그렇지 않다 — 두 인자 모두
    렌더 진입부에서 정규화된다는 것을 함께 본다.
    """
    file = render_export(
        export_format=ExportFormat.DOCX, title="안내\x0b문", body="본문\x00입니다.\n\n둘째\x0c 문단"
    )

    assert _paragraphs(file.content) == ["안내문", "본문입니다.", "둘째 문단"]


# --- txt ----------------------------------------------------------------------


def test_txt에도_제어문자가_남지_않는다() -> None:
    """정규화는 형식 분기 앞에서 한다 — docx와 txt가 같은 본문을 내놓아야 한다."""
    file = render_export(
        export_format=ExportFormat.TXT, title="안내문", body="본문\x00입니다.\n\n둘째\x0c 문단"
    )

    assert file.content.decode("utf-8") == "본문입니다.\n\n둘째 문단"


def test_복원한_원문에_섞인_제어문자도_걷어낸다() -> None:
    """AI 초안·마스킹 원문은 저장 경로에서 정규화되지 않는다 — 여기가 유일한 방어다."""
    body = restore_placeholders(
        "번호는 [[주민등록번호1]]이에요.", {"[[주민등록번호1]]": "900101-\x001234567"}
    )

    file = render_export(export_format=ExportFormat.TXT, title="안내문", body=body)

    assert file.content.decode("utf-8") == "번호는 900101-1234567이에요."


def test_txt는_BOM_없는_UTF_8이다() -> None:
    file = render_export(export_format=ExportFormat.TXT, title="안내문", body="본문입니다.")

    assert not file.content.startswith(b"\xef\xbb\xbf")
    assert file.content.decode("utf-8") == "본문입니다."
    assert file.filename == "안내문.txt"
    assert file.media_type == "text/plain; charset=utf-8"


# --- hwpx ---------------------------------------------------------------------


def _hwpx_text(content: bytes) -> str:
    """생성한 hwpx를 **우리 ingest 추출기로** 다시 읽는다.

    한컴 오피스가 없어 실제 열람은 확인할 수 없으므로(export.py 참고), 왕복 검증의
    기준은 읽기 경로다 — 우리가 이미 파싱하는 OWPML 구조로 조립됐는지를 본다.
    """
    return extract_text("내려받은 파일.hwpx", content)


def test_hwpx는_우리_추출기로_다시_열린다() -> None:
    """왕복: 생성 → 추출에서 제목·문단이 그대로 나와야 한다."""
    file = render_export(
        export_format=ExportFormat.HWPX,
        title="재난지원금 안내",
        body="첫 문단입니다.\n\n둘째 문단입니다.",
    )

    assert _hwpx_text(file.content) == "재난지원금 안내\n첫 문단입니다.\n둘째 문단입니다."
    assert file.filename == "재난지원금 안내.hwpx"
    assert file.media_type == "application/hwp+zip"


def test_hwpx는_mimetype을_무압축_첫_항목으로_담는다() -> None:
    """OWPML 패키지 규약(ODF와 같다) — 매직 바이트로 형식을 알아보는 도구를 위한 자리다."""
    file = render_export(export_format=ExportFormat.HWPX, title="안내문", body="본문입니다.")

    with zipfile.ZipFile(io.BytesIO(file.content)) as archive:
        first = archive.infolist()[0]
        assert first.filename == "mimetype"
        assert first.compress_type == zipfile.ZIP_STORED
        assert archive.read("mimetype") == b"application/hwp+zip"
        # 본문 구역과 패키지 뼈대가 함께 들어 있어야 한다.
        assert {"version.xml", "META-INF/container.xml", "Contents/header.xml"} <= set(
            archive.namelist()
        )


def test_hwpx에_XML_특수문자를_그대로_담는다() -> None:
    """&·<·>를 이스케이프하지 않으면 XML이 깨져 추출 자체가 실패한다."""
    body = '접수는 <온라인> & "우편" 둘 다 됩니다.\n\n기간 > 3월 2일'
    file = render_export(export_format=ExportFormat.HWPX, title="접수 안내 & 유의사항", body=body)

    assert _hwpx_text(file.content) == (
        '접수 안내 & 유의사항\n접수는 <온라인> & "우편" 둘 다 됩니다.\n기간 > 3월 2일'
    )


def test_hwpx에도_제어문자가_남지_않는다() -> None:
    """정규화는 형식 분기 앞에서 한다 — 세 형식이 같은 본문을 내놓아야 한다."""
    file = render_export(
        export_format=ExportFormat.HWPX, title="안내\x0b문", body="본문\x00입니다.\n\n둘째\x0c 문단"
    )

    assert _hwpx_text(file.content) == "안내문\n본문입니다.\n둘째 문단"


def test_hwpx는_빈_문단을_만들지_않는다() -> None:
    file = render_export(
        export_format=ExportFormat.HWPX, title="안내문", body="\n\n앞\n\n\n\n뒤\n\n"
    )

    assert _hwpx_text(file.content) == "안내문\n앞\n뒤"


def test_hwpx는_문단_안의_줄바꿈도_문단으로_나눈다() -> None:
    """hwpx 문단에는 강제 줄바꿈 대신 문단을 하나 더 쓴다 — 본문 줄 구성은 그대로다."""
    file = render_export(
        export_format=ExportFormat.HWPX, title="안내문", body="첫째 줄\n둘째 줄\n\n다음 문단"
    )

    assert _hwpx_text(file.content) == "안내문\n첫째 줄\n둘째 줄\n다음 문단"


def test_본문이_비어도_hwpx가_열린다() -> None:
    """제목만 남은 산출물도 깨진 패키지가 아니어야 한다."""
    file = render_export(export_format=ExportFormat.HWPX, title="안내문", body="   \n\n  ")

    assert _hwpx_text(file.content) == "안내문"


# --- 파일명 -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("재난지원금 안내", "재난지원금 안내.txt"),
        # 경로 구분자·제어문자는 제목에 들어올 수 있다(첫 줄에서 유도한 이름).
        # 앞뒤 점도 걷어낸다 — 숨김 파일(.name)이나 윈도우 금지 이름(name.)이 된다.
        ("../../etc/passwd", "etc passwd.txt"),
        ('따옴표"와\\역슬래시', "따옴표 와 역슬래시.txt"),
        ("줄바꿈\r\n섞임", "줄바꿈 섞임.txt"),
        # 전부 걸러지면 이름이 사라진다 — 대체 이름을 쓴다.
        ("///", "쉬운 글.txt"),
        ("   ", "쉬운 글.txt"),
    ],
)
def test_제목에서_안전한_파일명을_만든다(title: str, expected: str) -> None:
    assert export_filename(title, ExportFormat.TXT) == expected


def test_아주_긴_제목은_잘린다() -> None:
    """파일 시스템 이름 한계(대개 255바이트)를 넘기면 저장 자체가 실패한다."""
    name = export_filename("가" * 300, ExportFormat.DOCX)

    assert len(name.encode("utf-8")) < 255
    assert name.endswith(".docx")


# --- Content-Disposition -------------------------------------------------------


def test_한글_파일명은_RFC_5987로_인코딩한다() -> None:
    """HTTP 헤더는 latin-1만 담는다 — 한글을 그대로 넣으면 인코딩 오류가 난다."""
    value = content_disposition("재난지원금 안내.docx")

    assert value.startswith("attachment;")
    assert f"filename*=UTF-8''{quote('재난지원금 안내.docx', safe='')}" in value
    # ASCII만 아는 클라이언트를 위한 대체 이름이 함께 있어야 한다.
    assert 'filename="easy-read.docx"' in value
    # 헤더에 실을 수 있는 문자만 남는다(latin-1 인코딩이 통과해야 한다).
    value.encode("latin-1")
