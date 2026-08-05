"""업로드 파일 → 평문 텍스트 추출.

지원 형식은 docx·pdf·hwpx다. 구버전 hwp(바이너리 포맷)는 이번 범위 밖이라
`UnsupportedFormatError`로 명시적으로 거부한다 — 베스트에포트 파싱은 후속 과제.

**신뢰할 수 없는 입력을 다루는 모듈이라는 점**이 설계를 지배한다.

- 예외 메시지에 파일명·본문 조각을 넣지 않는다. 파일명 자체가 개인정보일 수 있고
  (예: `홍길동_주민등록등본.pdf`), 라이브러리 예외에는 임시 경로나 원문 일부가
  섞여 들어온다. 그래서 라이브러리 예외는 전부 `from None`으로 끊고 형식명만 담은
  고정 메시지로 바꾼다. `__suppress_context__`가 서므로 traceback 출력에도 원본
  메시지가 따라붙지 않는다(객체 참조는 `__context__`에 남지만 출력되지 않는다).
- 이 모듈은 로그를 남기지 않는다. 남길 만한 것이 파일명·본문뿐이라 전부 금지 대상이다.
- 압축 폭탄(docx·hwpx 둘 다 zip 컨테이너다)과 XML 엔티티 폭탄을 상한·거부로 막는다.
  아래 `_ensure_zip_not_bomb`·`_read_hwpx_sections`·`_hwpx_blocks` 주석 참고.

추출 결과의 모양은 형식과 무관하게 하나로 맞춘다: **공백뿐인 줄 없이, 문단·표 셀·
페이지가 개행 하나로 이어진 텍스트**. 후속 단계(마스킹·프롬프트)가 형식별 분기 없이
같은 입력을 받게 하기 위해서다.
"""

import io
import re
import zipfile
from collections.abc import Callable, Iterable, Iterator
from pathlib import PurePosixPath
from xml.etree import ElementTree

import docx
from docx.document import Document
from pypdf import PdfReader

from app.exceptions import DocumentExtractionError, UnsupportedFormatError

# 업로드 크기 상한. 실제 검사는 API 계층에서 한다 (여기서는 기준값만 정의).
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

# zip 컨테이너(docx·hwpx)를 풀었을 때 허용하는 총 크기 상한. 압축 폭탄 방어 —
# deflate는 반복 바이트를 1000:1 이상으로 줄이므로, 상한 안의 업로드도 수 GB로 부풀어
# 메모리를 고갈시킬 수 있다. 업로드 상한의 5배: 마크업이 본문보다 훨씬 긴 오피스 포맷의
# 특성을 감안한 여유값이다.
_MAX_UNCOMPRESSED_BYTES = 5 * MAX_UPLOAD_BYTES

# OWPML 패키지에서 본문을 담는 항목. 번호가 구역 순서다.
_SECTION_NAME = re.compile(r"^Contents/section(\d+)\.xml$")

# ElementTree는 내부 DTD의 엔티티를 그대로 펼친다("billion laughs"). OWPML 문서는
# DTD를 쓰지 않으므로 DOCTYPE 선언이 보이면 파싱 전에 거부한다.
_DOCTYPE = b"<!DOCTYPE"


def _broken(format_name: str) -> DocumentExtractionError:
    """손상 파일 예외. 형식명 외에는 아무것도 담지 않는다 (파일명·본문 유출 방지)."""
    return DocumentExtractionError(f"{format_name} 파일을 읽을 수 없습니다 (파일이 손상되었습니다)")


def _join_blocks(blocks: Iterable[str]) -> str:
    """블록(문단·표 셀·페이지) 텍스트를 정규화해 개행 하나로 잇는다.

    줄 단위로 좌우 공백을 털고 빈 줄을 버린다. 워드·PDF는 빈 문단과 줄 끝 공백을
    잔뜩 만들어내는데, 그대로 두면 프롬프트 토큰만 먹고 문단 경계도 흐려진다.
    """
    lines = [
        stripped for block in blocks for line in block.splitlines() if (stripped := line.strip())
    ]
    return "\n".join(lines)


def _docx_blocks(document: Document) -> Iterator[str]:
    """docx 본문 텍스트를 문단 → 표 셀 순서로 흘린다.

    한계: 표는 본문 뒤에 몰아서 붙는다(문서 내 원래 위치가 아니다). 또 병합 셀은
    python-docx가 행마다 같은 셀을 돌려주므로 텍스트가 중복될 수 있다. 쉬운 글 변환
    입력으로는 감수 가능한 손실이라 보고 단순한 순회를 택했다.
    """
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text


def _ensure_zip_not_bomb(data: bytes, format_name: str) -> None:
    """zip 항목의 해제 후 총 크기가 상한 안인지 파싱 전에 확인한다.

    python-docx는 압축 해제를 스스로 하므로 우리가 읽는 양을 제어할 수 없다. 대신 zip
    헤더에 선언된 크기를 미리 더해 본다 — Python zipfile은 항목을 선언 크기까지만 풀고
    초과분은 CRC·크기 검증에서 BadZipFile로 막으므로, 선언값은 위조해도 실제 해제량을
    늘릴 수 없는 신뢰 가능한 상한이다.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except Exception:
        raise _broken(format_name) from None
    if total > _MAX_UNCOMPRESSED_BYTES:
        raise DocumentExtractionError(f"{format_name} 파일이 너무 큽니다")


def _extract_docx(data: bytes) -> str:
    """python-docx로 문단과 표 셀 텍스트를 뽑는다."""
    _ensure_zip_not_bomb(data, "docx")
    try:
        document = docx.Document(io.BytesIO(data))
        # 지연 파싱 구간을 try 안에서 소진한다 — 밖에서 터지면 예외가 그대로 새어 나간다.
        blocks = list(_docx_blocks(document))
    except Exception:
        raise _broken("docx") from None
    return _join_blocks(blocks)


def _extract_pdf(data: bytes) -> str:
    """pypdf로 페이지별 텍스트를 뽑아 잇는다."""
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() for page in reader.pages]
    except Exception:
        raise _broken("pdf") from None
    text = _join_blocks(pages)
    if not text:
        # 텍스트 레이어가 없는 스캔 PDF. 손상과 구분되는 상태이므로 안내를 달리한다.
        raise DocumentExtractionError("텍스트를 추출할 수 없습니다 (스캔 PDF는 지원 예정)")
    return text


def _read_hwpx_sections(data: bytes) -> list[bytes]:
    """hwpx(zip)에서 구역 XML을 번호 순서로 읽는다.

    zip 기록 순서는 믿지 않고 `section{N}.xml`의 N으로 정렬한다. 크기도 헤더 선언값이
    아니라 실제 읽은 바이트로 세고, 남은 예산까지만 읽어 해제량이 상한을 넘는 순간
    끊는다 — 다 읽고 나서 확인하면 확인 시점에 이미 메모리가 차 있다.
    """
    budget = _MAX_UNCOMPRESSED_BYTES
    sections: list[bytes] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = sorted(
                (int(match.group(1)), match.group(0))
                for name in archive.namelist()
                if (match := _SECTION_NAME.match(name))
            )
            for _, name in names:
                with archive.open(name) as member:
                    raw = member.read(budget + 1)
                sections.append(raw)
                budget -= len(raw)
                if budget < 0:
                    break
    except Exception:
        raise _broken("hwpx") from None
    if budget < 0:
        raise DocumentExtractionError("hwpx 파일이 너무 큽니다")
    return sections


def _hwpx_blocks(section: bytes) -> list[str]:
    """OWPML 구역 XML에서 문단 단위 텍스트를 뽑는다.

    구조(KS X 6101): 구역 `<hs:sec>` > 문단 `<hp:p>` > 글자 조각 `<hp:run>/<hp:t>`.
    한 문단이 서식 때문에 여러 `<hp:t>`로 쪼개지므로 조각은 구분자 없이 잇고,
    `<hp:p>`를 만날 때만 블록을 끊는다. 네임스페이스 URI가 버전마다 달라 접두사가
    아니라 로컬 이름으로 판별한다.

    한계: 표 셀 안의 중첩 문단은 바깥 문단과 별도 블록이 되며(원하는 동작), 바깥 문단이
    중첩 문단 뒤에 텍스트를 더 가지면 그 텍스트는 중첩 블록에 붙는다. 실제 hwpx에서
    표를 담은 문단은 자체 텍스트가 없어 문제되지 않는다.
    """
    if _DOCTYPE in section:
        raise DocumentExtractionError("hwpx 파일을 읽을 수 없습니다 (허용되지 않는 XML 선언)")
    try:
        root = ElementTree.fromstring(section)
    except Exception:
        raise _broken("hwpx") from None

    blocks: list[str] = []
    current: list[str] = []
    for element in root.iter():
        local_name = element.tag.rpartition("}")[2]
        if local_name == "p":
            blocks.append("".join(current))
            current = []
        elif local_name == "t":
            current.append("".join(element.itertext()))
    blocks.append("".join(current))
    return blocks


def _extract_hwpx(data: bytes) -> str:
    """zip + ElementTree로 OWPML 본문 텍스트를 뽑는다."""
    sections = _read_hwpx_sections(data)
    if not sections:
        # 구역이 하나도 없으면 hwpx 패키지가 아니거나 껍데기다.
        raise _broken("hwpx")
    return _join_blocks(block for section in sections for block in _hwpx_blocks(section))


_EXTRACTORS: dict[str, Callable[[bytes], str]] = {
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
    ".hwpx": _extract_hwpx,
}


def extract_text(filename: str, data: bytes) -> str:
    """업로드 파일에서 본문 텍스트를 뽑는다.

    확장자로만 형식을 판별한다(내용 스니핑 없음). 대소문자는 가리지 않는다.

    Raises:
        UnsupportedFormatError: 지원하지 않는 확장자 (구버전 hwp 포함).
        DocumentExtractionError: 파일이 손상됐거나 추출 가능한 텍스트가 없다.
    """
    extension = PurePosixPath(filename).suffix.lower()
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise UnsupportedFormatError("지원 형식: docx, pdf, hwpx")
    return extractor(data)
